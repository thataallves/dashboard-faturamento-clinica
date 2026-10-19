"""
Dados fictícios inspirados na rotina de uma clínica odontológica com HOF
(Clínica Sorriso Pleno, São Gonçalo - RJ). Nomes, pacientes e atendimentos
são gerados por script.
Projeto 1 (Excel -> Power BI).
Faturamento mensal calibrado entre ~R$ 50 mil e ~R$ 80 mil.
"""
import numpy as np
import pandas as pd

rng = np.random.default_rng(11)

dentistas = pd.DataFrame({
    "id_dentista": [1, 2, 3, 4],
    "dentista": ["Dra. Helena Prado", "Dra. Camila Duarte", "Dra. Marina Teles", "Dra. Thaïs Alves"],
    "especialidade": ["Ortodontia", "Implantodontia", "Endodontia", "Clínica Geral e HOF"],
})

# nome, categoria, valor_tabela (R$), peso, especialidade responsável
procs = [
    ("Consulta / Avaliação",        "Diagnóstico",  80, 14, None),
    ("Profilaxia (Limpeza)",        "Preventivo",  120, 14, "Clínica Geral e HOF"),
    ("Restauração em Resina",       "Dentística",  160, 14, "Clínica Geral e HOF"),
    ("Raspagem Periodontal",        "Periodontia", 200,  5, "Clínica Geral e HOF"),
    ("Clareamento",                 "Estética",    500,  5, "Clínica Geral e HOF"),
    ("Extração Simples",            "Cirurgia",    150,  5, "Clínica Geral e HOF"),
    ("Faceta em Resina",            "Estética",    400,  1.5, "Clínica Geral e HOF"),
    ("Toxina Botulínica",           "HOF",         900,  4, "Clínica Geral e HOF"),
    ("Preenchimento Labial",        "HOF",         850,  3, "Clínica Geral e HOF"),
    ("Preenchimento de Olheiras",   "HOF",         900,  1.2, "Clínica Geral e HOF"),
    ("Bioestimulador de Colágeno",  "HOF",        1500,  1, "Clínica Geral e HOF"),
    # --- outras especialidades da clínica (outras dentistas) ---
    ("Manutenção Ortodôntica",      "Ortodontia",  120, 18, "Ortodontia"),
    ("Instalação de Aparelho",      "Ortodontia",  900,  2.5, "Ortodontia"),
    ("Tratamento de Canal",         "Endodontia",  650,  8, "Endodontia"),
    ("Implante Dentário",           "Implantodontia", 2200, 1.5, "Implantodontia"),
    ("Extração de Siso",            "Cirurgia",    350,  3, "Implantodontia"),
]
procedimentos = pd.DataFrame(procs, columns=["procedimento", "categoria", "valor_tabela", "peso", "especialidade"])
procedimentos.insert(0, "id_procedimento", range(1, len(procedimentos) + 1))
# equilibra a clínica: clínica geral + HOF não concentra tudo (orto/endo/implante são fortes)
procedimentos.loc[procedimentos["especialidade"] == "Clínica Geral e HOF", "peso"] *= 0.55

formas = ["Pix", "Cartão de crédito", "Cartão de débito", "Dinheiro", "Convênio"]
p_formas = [0.38, 0.40, 0.10, 0.07, 0.05]   # HOF é quase todo particular

metas = {
 2024: [50000, 52000, 58000, 62000, 60000, 62000, 56000, 64000, 66000, 70000, 72000, 60000],
 2025: [58000, 60000, 66000, 70000, 68000, 70000, 64000, 72000, 74000, 78000, 80000, 66000],
}
p_proc = procedimentos["peso"] / procedimentos["peso"].sum()

linhas = []
for ano, lista in metas.items():
    for mes, meta in enumerate(lista, start=1):
        meta = meta * rng.uniform(0.985, 1.015)
        fim = pd.Timestamp(ano, mes, 1) + pd.offsets.MonthEnd(0)
        dias = pd.bdate_range(f"{ano}-{mes:02d}-01", fim)
        dias = dias.append(pd.date_range(f"{ano}-{mes:02d}-01", fim, freq="W-SAT")[::2]).sort_values()
        w = np.array([0.5 if d.dayofweek == 5 else 1.0 for d in dias]); w /= w.sum()
        bloco, total = [], 0.0
        while total < meta:
            d = dias[rng.choice(len(dias), p=w)]
            proc = procedimentos.iloc[rng.choice(len(procedimentos), p=p_proc)]
            if pd.isna(proc["especialidade"]):
                cand = dentistas
            else:
                cand = dentistas[dentistas["especialidade"] == proc["especialidade"]]
            dent = cand.iloc[rng.choice(len(cand), p=[0.6, 0.4] if len(cand) == 2 else None)]
            desc = rng.choice([0, 0.05, 0.10], p=[0.80, 0.12, 0.08])
            valor = proc["valor_tabela"] * (1 - desc) * rng.uniform(0.97, 1.03)
            total += valor
            bloco.append([d, int(rng.integers(1, 700)), int(dent["id_dentista"]), int(proc["id_procedimento"]),
                          rng.choice(formas, p=p_formas), desc, valor])
        fator = meta / total
        for l in bloco:
            l[6] = round(l[6] * fator, 2)
        linhas += bloco

atend = pd.DataFrame(linhas, columns=["data", "id_paciente", "id_dentista", "id_procedimento",
                                      "forma_pagamento", "desconto_pct", "valor_cobrado"])
atend = atend.sort_values("data").reset_index(drop=True)
atend.insert(0, "id_atendimento", range(1, len(atend) + 1))

# "Sujeira" controlada para treinar o Power Query
idx = rng.choice(atend.index, size=int(len(atend) * 0.012), replace=False)
atend.loc[idx, "forma_pagamento"] = None
atend = pd.concat([atend, atend.sample(n=15, random_state=1)]).sort_values("id_atendimento").reset_index(drop=True)
idx = rng.choice(atend.index, size=25, replace=False)
atend.loc[idx, "forma_pagamento"] = atend.loc[idx, "forma_pagamento"].replace(
    {"Pix": "PIX ", "Cartão de crédito": "cartao credito"})

saida = "/home/claude/projeto1/Clinica_Sorriso_Pleno_Atendimentos.xlsx"
with pd.ExcelWriter(saida, engine="openpyxl", datetime_format="DD/MM/YYYY") as w:
    atend.to_excel(w, sheet_name="Atendimentos", index=False)
    dentistas.to_excel(w, sheet_name="Dentistas", index=False)
    procedimentos.drop(columns=["peso", "especialidade"]).to_excel(w, sheet_name="Procedimentos", index=False)

from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
wb = load_workbook(saida)
for ws, nome in zip(wb.worksheets, ["tbAtendimentos", "tbDentistas", "tbProcedimentos"]):
    t = Table(displayName=nome, ref=f"A1:{ws.cell(row=ws.max_row, column=ws.max_column).coordinate}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(t)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(14, len(str(col[0].value)) + 4)
    ws.freeze_panes = "A2"
wb.save(saida)

base = atend.drop_duplicates("id_atendimento").merge(procedimentos[["id_procedimento", "categoria"]], on="id_procedimento")
m = base.groupby(base["data"].dt.to_period("M"))["valor_cobrado"].agg(["sum", "count"])
m["ticket"] = m["sum"] / m["count"]
print(m.round(0).iloc[[0, 1, 2, 11, 12, 22, 23]].to_string())
print("Total: R$ {:,.0f} | Atendimentos: {}".format(base["valor_cobrado"].sum(), len(base)))
print((base.groupby("categoria")["valor_cobrado"].sum() / base["valor_cobrado"].sum()).round(2))
